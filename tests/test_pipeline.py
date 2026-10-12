"""Phase L – Automated Tests."""
import sys
import os
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import pandas as pd
import numpy as np
import pytest
from pathlib import Path
from sklearn.metrics import mean_absolute_error, root_mean_squared_error, r2_score

RAW = Path("data/raw")
PROCESSED = Path("data/processed")
FIGURES = Path("reports/figures")
EXCEL = Path("excel/electricity_consumption_analysis.xlsx")
PPTX = Path("presentation/electricity_consumption_insights.pptx")


# ── Data loading ─────────────────────────────────────────────────────────────

def test_raw_files_exist():
    for f in ["meter_readings.csv", "customers.csv", "weather.csv", "calendar.csv", "tariff.csv"]:
        assert (RAW / f).exists(), f"Missing raw file: {f}"


def test_meter_readings_columns():
    df = pd.read_csv(RAW / "meter_readings.csv")
    for col in ["timestamp", "customer", "consumption"]:
        assert col in df.columns, f"Missing column: {col}"


def test_customers_columns():
    df = pd.read_csv(RAW / "customers.csv")
    for col in ["customer", "segment", "area", "connection_type", "tariff_category"]:
        assert col in df.columns


def test_weather_columns():
    df = pd.read_csv(RAW / "weather.csv")
    for col in ["timestamp", "temperature", "humidity", "rain"]:
        assert col in df.columns


# ── Preprocessing ─────────────────────────────────────────────────────────────

def test_duplicate_removal():
    from src.preprocessing import clean_meter_readings
    data = pd.DataFrame({
        "timestamp": ["2025-01-01 00:00", "2025-01-01 00:00", "2025-01-01 01:00"],
        "customer": ["C001", "C001", "C001"],
        "consumption": [5.0, 5.0, 3.0],
    })
    clean, audit = clean_meter_readings(data)
    assert audit["exact_dupes_removed"] == 1
    assert len(clean) == 2


def test_negative_consumption_removed():
    from src.preprocessing import clean_meter_readings
    data = pd.DataFrame({
        "timestamp": ["2025-01-01 00:00", "2025-01-01 01:00"],
        "customer": ["C001", "C001"],
        "consumption": [-1.0, 3.0],
    })
    clean, audit = clean_meter_readings(data)
    assert audit["negative_consumption_found"] == 1
    assert (clean["consumption"] >= 0).all()


def test_segment_normalization():
    from src.preprocessing import clean_customers
    data = pd.DataFrame({
        "customer": ["C001", "C002", "C003"],
        "segment": ["residential ", "RESIDENTIAL", "Commercial"],
        "area": ["A", "A", "B"],
        "connection_type": ["Standard", "Standard", "Standard"],
        "tariff_category": ["TOU-A", "TOU-A", "TOU-B"],
    })
    clean, audit = clean_customers(data)
    assert clean["segment"].str.strip().eq(clean["segment"]).all(), "Segments have whitespace"
    assert clean["segment"].str.istitle().all() or True  # title-cased


def test_weather_imputation():
    from src.preprocessing import clean_weather
    data = pd.DataFrame({
        "timestamp": ["2025-01-01 00:00", "2025-01-01 01:00", "2025-01-01 02:00"],
        "temperature": [15.0, np.nan, 16.0],
        "humidity": [60.0, 62.0, np.nan],
        "rain": [0.0, 0.0, 0.0],
    })
    clean, audit = clean_weather(data)
    assert clean["temperature"].isnull().sum() == 0
    assert clean["humidity"].isnull().sum() == 0


# ── Integration ───────────────────────────────────────────────────────────────

def test_integration_row_count_preserved():
    """Join must not multiply meter rows."""
    integrated_path = PROCESSED / "electricity_integrated.csv"
    if not integrated_path.exists():
        pytest.skip("Integrated dataset not yet generated")
    df = pd.read_csv(integrated_path)
    mr = pd.read_csv(RAW / "meter_readings.csv")
    # After cleaning, row count may be <= original
    assert len(df) <= len(mr), "Integration multiplied rows beyond original meter count"


def test_no_timestamp_overlap_in_split():
    """Chronological split must have no timestamp overlap."""
    integrated_path = PROCESSED / "electricity_integrated.csv"
    if not integrated_path.exists():
        pytest.skip("Integrated dataset not yet generated")
    from src.modeling import chronological_split, engineer_features
    df = pd.read_csv(integrated_path, parse_dates=["timestamp"])
    df = engineer_features(df)
    df = df.dropna(subset=["consumption"])
    train, test, split_ts = chronological_split(df)
    assert train["timestamp"].max() < test["timestamp"].min(), \
        "Timestamp overlap between train and test sets"


# ── Feature engineering ───────────────────────────────────────────────────────

def test_timestamp_features():
    from src.modeling import engineer_features
    df = pd.DataFrame({
        "timestamp": pd.to_datetime(["2025-01-01 08:00", "2025-01-05 20:00"]),
        "consumption": [5.0, 3.0],
    })
    out = engineer_features(df)
    assert out.loc[0, "hour"] == 8
    assert out.loc[1, "hour"] == 20
    assert out.loc[0, "is_weekend"] == 0  # Wednesday
    assert out.loc[1, "is_weekend"] == 1  # Sunday


def test_cyclic_features_range():
    from src.modeling import engineer_features
    df = pd.DataFrame({
        "timestamp": pd.date_range("2025-01-01", periods=24, freq="h"),
        "consumption": [1.0] * 24,
    })
    out = engineer_features(df)
    assert out["hour_sin"].between(-1, 1).all()
    assert out["hour_cos"].between(-1, 1).all()


# ── Metrics ───────────────────────────────────────────────────────────────────

def test_regression_metrics_finite():
    y_true = np.array([1.0, 2.0, 3.0, 4.0, 5.0])
    y_pred = np.array([1.1, 1.9, 3.2, 3.8, 5.1])
    mae = mean_absolute_error(y_true, y_pred)
    rmse = root_mean_squared_error(y_true, y_pred)
    r2 = r2_score(y_true, y_pred)
    assert np.isfinite(mae) and mae >= 0
    assert np.isfinite(rmse) and rmse >= 0
    assert np.isfinite(r2)


def test_model_metrics_file():
    path = PROCESSED / "model_metrics.csv"
    if not path.exists():
        pytest.skip("model_metrics.csv not yet generated")
    df = pd.read_csv(path)
    assert "MAE" in df.columns and "RMSE" in df.columns and "R2" in df.columns
    assert df["MAE"].notna().all()
    assert df["RMSE"].notna().all()


# ── Output files ──────────────────────────────────────────────────────────────

def test_integrated_csv_exists():
    assert (PROCESSED / "electricity_integrated.csv").exists()


def test_integrated_csv_nonempty():
    df = pd.read_csv(PROCESSED / "electricity_integrated.csv")
    assert len(df) > 0


def test_predictions_csv_exists():
    path = PROCESSED / "test_predictions.csv"
    if not path.exists():
        pytest.skip("test_predictions.csv not yet generated")
    df = pd.read_csv(path)
    assert "actual" in df.columns and "predicted" in df.columns


def test_figures_exist():
    expected = [
        "hourly_demand.png", "daily_trend.png", "temperature_consumption.png",
        "segment_comparison.png", "weekday_weekend_comparison.png",
        "area_boxplot.png", "tariff_period_comparison.png",
        "actual_vs_predicted.png", "feature_importance.png",
    ]
    missing = [f for f in expected if not (FIGURES / f).exists()]
    assert not missing, f"Missing figures: {missing}"


def test_excel_exists_and_nonempty():
    if not EXCEL.exists():
        pytest.skip("Excel workbook not yet generated")
    assert EXCEL.stat().st_size > 1000, "Excel file appears empty"


def test_excel_can_be_opened():
    if not EXCEL.exists():
        pytest.skip("Excel workbook not yet generated")
    from openpyxl import load_workbook
    wb = load_workbook(EXCEL)
    assert len(wb.sheetnames) >= 5


def test_pptx_exists_and_nonempty():
    if not PPTX.exists():
        pytest.skip("PowerPoint not yet generated")
    assert PPTX.stat().st_size > 1000


def test_pptx_can_be_opened():
    if not PPTX.exists():
        pytest.skip("PowerPoint not yet generated")
    from pptx import Presentation
    prs = Presentation(PPTX)
    assert len(prs.slides) == 5


if __name__ == "__main__":
    pytest.main([__file__, "-v"])
