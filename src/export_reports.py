"""Phase G & I – Excel Workbook and PowerPoint Presentation."""
import pandas as pd
import numpy as np
from pathlib import Path
import textwrap

PROCESSED = Path("data/processed")
EXCEL_DIR = Path("excel")
PPTX_DIR = Path("presentation")
FIGURES = Path("reports/figures")

DISCLAIMER = "SYNTHETIC DATA — Not official GridSense data. For practice purposes only."


# ── Excel ────────────────────────────────────────────────────────────────────

def _autofit(ws, min_width=10, max_width=40):
    from openpyxl.utils import get_column_letter
    for col_cells in ws.columns:
        length = max(
            len(str(cell.value)) if cell.value is not None else 0
            for cell in col_cells
        )
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = \
            min(max(length + 2, min_width), max_width)


def _freeze(ws):
    ws.freeze_panes = "A2"


def _write_df(ws, df: pd.DataFrame, start_row=1):
    from openpyxl.styles import Font, PatternFill, Alignment
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF")
    for col_idx, col_name in enumerate(df.columns, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=col_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    for row_idx, row in enumerate(df.itertuples(index=False), start_row + 1):
        for col_idx, value in enumerate(row, 1):
            ws.cell(row=row_idx, column=col_idx, value=value)


def create_excel(eda_results: dict):
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Font, Alignment

    EXCEL_DIR.mkdir(exist_ok=True)
    wb = Workbook()

    # ── README ──
    ws = wb.active
    ws.title = "README"
    ws["A1"] = "Electricity Consumption Insight — Analysis Workbook"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = DISCLAIMER
    ws["A2"].font = Font(italic=True, color="FF0000")
    ws["A4"] = "Sheets:"
    sheets_info = [
        ("Cleaned_Data", "Sample of the cleaned integrated dataset"),
        ("Hourly_Demand", "Average aggregate demand by hour of day"),
        ("Weather_Analysis", "Weather correlation summary"),
        ("Segment_Analysis", "Consumption statistics by customer segment"),
        ("Weekend_Analysis", "Weekday vs weekend comparison"),
        ("Area_Analysis", "Consumption statistics by area"),
        ("Tariff_Analysis", "Consumption by tariff period"),
        ("Model_Metrics", "Baseline and Random Forest evaluation metrics"),
        ("Summary_Dashboard", "Key findings summary"),
        ("PivotTable_Instructions", "How to create native PivotTables in Excel"),
    ]
    for i, (name, desc) in enumerate(sheets_info, 5):
        ws[f"A{i}"] = f"  {name}: {desc}"
    ws.column_dimensions["A"].width = 70

    # ── Cleaned_Data ──
    df = pd.read_csv(PROCESSED / "electricity_integrated.csv")
    ws2 = wb.create_sheet("Cleaned_Data")
    sample = df.head(5000).copy()
    # Format timestamp
    if "timestamp" in sample.columns:
        sample["timestamp"] = sample["timestamp"].astype(str)
    if "date" in sample.columns:
        sample["date"] = sample["date"].astype(str)
    _write_df(ws2, sample)
    _freeze(ws2)
    _autofit(ws2)
    ws2["A1"].comment = None

    # ── Hourly_Demand ──
    hourly = eda_results["hourly"]
    ws3 = wb.create_sheet("Hourly_Demand")
    _write_df(ws3, hourly.round(4))
    _freeze(ws3)
    _autofit(ws3)
    # Embed chart image if available
    img_path = FIGURES / "hourly_demand.png"
    if img_path.exists():
        img = XLImage(str(img_path))
        img.anchor = "D2"
        ws3.add_image(img)

    # ── Weather_Analysis ──
    ws4 = wb.create_sheet("Weather_Analysis")
    corr = eda_results["corr"]
    corr_df = pd.DataFrame(list(corr.items()), columns=["Weather Variable", "Correlation with Demand"])
    _write_df(ws4, corr_df)
    _freeze(ws4)
    _autofit(ws4)
    img_path = FIGURES / "correlation_heatmap.png"
    if img_path.exists():
        img = XLImage(str(img_path))
        img.anchor = "D2"
        ws4.add_image(img)

    # ── Segment_Analysis ──
    ws5 = wb.create_sheet("Segment_Analysis")
    seg = eda_results["segment"]
    _write_df(ws5, seg.round(4))
    _freeze(ws5)
    _autofit(ws5)
    img_path = FIGURES / "segment_comparison.png"
    if img_path.exists():
        img = XLImage(str(img_path))
        img.anchor = "H2"
        ws5.add_image(img)

    # ── Weekend_Analysis ──
    ws6 = wb.create_sheet("Weekend_Analysis")
    ww_df = pd.DataFrame(eda_results["weekday_weekend"]["summary"])
    ww_df["type"] = ww_df["is_weekend"].map({False: "Weekday", True: "Weekend"})
    _write_df(ws6, ww_df[["type", "mean", "median", "total", "count"]].round(4))
    _freeze(ws6)
    _autofit(ws6)
    img_path = FIGURES / "weekday_weekend_comparison.png"
    if img_path.exists():
        img = XLImage(str(img_path))
        img.anchor = "H2"
        ws6.add_image(img)

    # ── Area_Analysis ──
    ws7 = wb.create_sheet("Area_Analysis")
    area = eda_results["area"]
    _write_df(ws7, area.round(4))
    _freeze(ws7)
    _autofit(ws7)
    img_path = FIGURES / "area_boxplot.png"
    if img_path.exists():
        img = XLImage(str(img_path))
        img.anchor = "J2"
        ws7.add_image(img)

    # ── Tariff_Analysis ──
    ws8 = wb.create_sheet("Tariff_Analysis")
    tariff = eda_results["tariff"]
    if not tariff.empty:
        _write_df(ws8, tariff.round(4))
        _freeze(ws8)
        _autofit(ws8)
    img_path = FIGURES / "tariff_period_comparison.png"
    if img_path.exists():
        img = XLImage(str(img_path))
        img.anchor = "H2"
        ws8.add_image(img)

    # ── Model_Metrics ──
    ws9 = wb.create_sheet("Model_Metrics")
    metrics_path = PROCESSED / "model_metrics.csv"
    if metrics_path.exists():
        metrics = pd.read_csv(metrics_path)
        _write_df(ws9, metrics)
        _freeze(ws9)
        _autofit(ws9)

    # ── Summary_Dashboard ──
    ws10 = wb.create_sheet("Summary_Dashboard")
    ws10["A1"] = "Electricity Consumption Insight — Summary Dashboard"
    ws10["A1"].font = Font(bold=True, size=14)
    ws10["A2"] = DISCLAIMER
    ws10["A2"].font = Font(italic=True, color="FF0000")
    stats = eda_results["stats"]
    dashboard_rows = [
        ("Total Consumption (kWh)", stats["total_consumption"]),
        ("Valid Observations", stats["valid_observations"]),
        ("Mean Consumption per Reading (kWh)", stats["mean"]),
        ("Median Consumption per Reading (kWh)", stats["median"]),
        ("Std Dev", stats["std"]),
        ("Max Consumption (kWh)", stats["max"]),
        ("Peak Demand Hour", f"{eda_results['peak_hour']}:00"),
        ("Low Demand Hour", f"{eda_results['low_hour']}:00"),
    ]
    for i, (label, value) in enumerate(dashboard_rows, 4):
        ws10[f"A{i}"] = label
        ws10[f"B{i}"] = value
    ws10.column_dimensions["A"].width = 40
    ws10.column_dimensions["B"].width = 20

    # ── PivotTable_Instructions ──
    ws11 = wb.create_sheet("PivotTable_Instructions")
    ws11["A1"] = "How to Create Native PivotTables in Excel"
    ws11["A1"].font = Font(bold=True, size=13)
    instructions = [
        "",
        "NOTE: The Cleaned_Data sheet contains pivot-ready data (up to 5,000 rows).",
        "For full data, use the file: data/processed/electricity_integrated.csv",
        "",
        "PIVOT 1 — Hourly Demand",
        "1. Select the Cleaned_Data sheet.",
        "2. Insert > PivotTable > New Sheet.",
        "3. Drag 'hour' to Rows.",
        "4. Drag 'consumption' to Values (set to Average).",
        "5. Insert > PivotChart > Line chart.",
        "",
        "PIVOT 2 — Segment Comparison",
        "1. Insert > PivotTable on Cleaned_Data.",
        "2. Drag 'segment' to Rows.",
        "3. Drag 'consumption' to Values (Average and Sum).",
        "4. Insert > PivotChart > Bar chart.",
        "",
        "PIVOT 3 — Tariff Period",
        "1. Insert > PivotTable on Cleaned_Data.",
        "2. Drag 'peak_period' to Rows.",
        "3. Drag 'consumption' to Values (Average).",
        "",
        "PIVOT 4 — Area Analysis",
        "1. Insert > PivotTable on Cleaned_Data.",
        "2. Drag 'area' to Rows.",
        "3. Drag 'consumption' to Values (Average, Sum, StdDev).",
        "",
        "These PivotTables are native Excel features and will be refreshable",
        "if you load the full CSV as a data source.",
    ]
    for i, line in enumerate(instructions, 2):
        ws11[f"A{i}"] = line
    ws11.column_dimensions["A"].width = 70

    path = EXCEL_DIR / "electricity_consumption_analysis.xlsx"
    wb.save(path)
    print(f"Excel workbook saved: {path} ({path.stat().st_size:,} bytes)")
    return path


# ── PowerPoint ───────────────────────────────────────────────────────────────

def create_pptx(eda_results: dict, model_results: dict):
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN

    PPTX_DIR.mkdir(exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    DARK_BLUE = RGBColor(0x1F, 0x4E, 0x79)
    MID_BLUE = RGBColor(0x2E, 0x75, 0xB6)
    WHITE = RGBColor(0xFF, 0xFF, 0xFF)
    RED = RGBColor(0xC0, 0x00, 0x00)

    blank_layout = prs.slide_layouts[6]  # blank

    def add_slide():
        return prs.slides.add_slide(blank_layout)

    def add_textbox(slide, text, left, top, width, height,
                    bold=False, size=18, color=None, wrap=True, align=PP_ALIGN.LEFT):
        from pptx.util import Inches, Pt
        txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        tf = txBox.text_frame
        tf.word_wrap = wrap
        p = tf.paragraphs[0]
        p.alignment = align
        run = p.add_run()
        run.text = text
        run.font.bold = bold
        run.font.size = Pt(size)
        if color:
            run.font.color.rgb = color
        return txBox

    def add_rect(slide, left, top, width, height, fill_color):
        from pptx.util import Inches
        shape = slide.shapes.add_shape(1, Inches(left), Inches(top), Inches(width), Inches(height))
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
        shape.line.fill.background()
        return shape

    def add_image_safe(slide, img_name, left, top, width):
        img_path = FIGURES / img_name
        if img_path.exists():
            slide.shapes.add_picture(str(img_path), Inches(left), Inches(top), width=Inches(width))

    stats = eda_results["stats"]
    peak_hour = eda_results["peak_hour"]
    low_hour = eda_results["low_hour"]
    corr = eda_results["corr"]
    seg = eda_results["segment"]
    top_seg = seg.loc[seg["mean"].idxmax(), "segment"]
    top_seg_mean = seg.loc[seg["mean"].idxmax(), "mean"]
    rf = model_results["rf"]
    baseline = model_results["baseline"]

    # ── Slide 1: Business Problem ──
    sl = add_slide()
    add_rect(sl, 0, 0, 13.33, 7.5, DARK_BLUE)
    add_textbox(sl, "Electricity Consumption Insight", 0.5, 0.4, 12, 1.2,
                bold=True, size=36, color=WHITE, align=PP_ALIGN.CENTER)
    add_textbox(sl, "GridSense Energy Services", 0.5, 1.6, 12, 0.6,
                size=22, color=RGBColor(0xBD, 0xD7, 0xEE), align=PP_ALIGN.CENTER)
    add_textbox(sl, "Business Question:", 0.8, 2.5, 11, 0.5, bold=True, size=18, color=WHITE)
    add_textbox(sl,
                "What temporal and environmental factors influence electricity consumption?",
                0.8, 3.0, 11.5, 0.8, size=20, color=RGBColor(0xFF, 0xC0, 0x00))
    objectives = (
        "• Identify high- and low-demand periods\n"
        "• Analyze weather–consumption relationships\n"
        "• Compare customer segments and geographical areas\n"
        "• Investigate weekday/weekend and tariff-period patterns\n"
        "• Train and evaluate a regression model"
    )
    add_textbox(sl, objectives, 0.8, 3.9, 11.5, 2.5, size=16, color=WHITE)
    add_textbox(sl, DISCLAIMER, 0.5, 7.0, 12, 0.4, size=10,
                color=RGBColor(0xFF, 0x80, 0x80), align=PP_ALIGN.CENTER)

    # ── Slide 2: Dataset and Preparation ──
    sl = add_slide()
    add_rect(sl, 0, 0, 13.33, 1.1, DARK_BLUE)
    add_textbox(sl, "Dataset and Preparation", 0.3, 0.15, 12, 0.8,
                bold=True, size=28, color=WHITE)
    add_textbox(sl, "Datasets", 0.4, 1.2, 4, 0.4, bold=True, size=16, color=DARK_BLUE)
    dataset_text = (
        "• meter_readings: 229,658 rows (timestamp, customer, consumption)\n"
        "• customers: 80 rows (segment, area, connection_type, tariff_category)\n"
        "• weather: 2,868 rows (temperature, humidity, rain)\n"
        "• calendar: 120 rows (date, weekday, holiday, season)\n"
        "• tariff: 72 rows (tariff_category, hour, peak_period, tariff_rate)\n"
        "• Coverage: 2025-01-01 to 2025-04-30 | 80 customers"
    )
    add_textbox(sl, dataset_text, 0.4, 1.65, 6, 2.8, size=14)
    add_textbox(sl, "Data Quality Issues & Fixes", 6.8, 1.2, 6, 0.4, bold=True, size=16, color=DARK_BLUE)
    quality_text = (
        "• Exact duplicate rows → removed\n"
        "• Conflicting customer+timestamp keys → kept higher consumption\n"
        "• Negative consumption values → removed (synthetic errors)\n"
        "• Inconsistent segment labels → normalized (strip + title-case)\n"
        "• Missing weather measurements → forward/back-filled\n"
        "• Missing consumption → excluded from ML training\n"
        "• All joins validated (many-to-one); row count preserved"
    )
    add_textbox(sl, quality_text, 6.8, 1.65, 6.2, 3.2, size=14)
    add_textbox(sl, DISCLAIMER, 0.5, 7.1, 12, 0.35, size=10,
                color=RED, align=PP_ALIGN.CENTER)

    # ── Slide 3: EDA Findings ──
    sl = add_slide()
    add_rect(sl, 0, 0, 13.33, 1.1, DARK_BLUE)
    add_textbox(sl, "EDA Findings", 0.3, 0.15, 12, 0.8, bold=True, size=28, color=WHITE)

    add_textbox(sl, f"Finding 1 — Peak Demand Hour: {peak_hour}:00",
                0.4, 1.2, 6.2, 0.5, bold=True, size=16, color=DARK_BLUE)
    f1_text = (
        f"Peak aggregate demand occurs at {peak_hour}:00.\n"
        f"Low demand occurs at {low_hour}:00.\n"
        f"Peak avg demand: {eda_results['hourly']['avg_aggregate_demand'].max():.1f} kWh\n"
        f"Low avg demand:  {eda_results['hourly']['avg_aggregate_demand'].min():.1f} kWh\n"
        f"Implication: Grid operators should schedule maintenance and demand-response\n"
        f"programs around the {peak_hour}:00 peak window."
    )
    add_textbox(sl, f1_text, 0.4, 1.75, 6.2, 2.5, size=13)
    add_image_safe(sl, "hourly_demand.png", 0.3, 4.1, 6.2)

    add_textbox(sl, f"Finding 2 — Segment Differences (Top: {top_seg})",
                6.8, 1.2, 6.2, 0.5, bold=True, size=16, color=DARK_BLUE)
    f2_text = (
        f"Highest mean consumption: {top_seg} ({top_seg_mean:.3f} kWh/reading)\n"
        f"Temp–demand correlation: r = {corr.get('temperature', 'N/A')}\n"
        f"Humidity–demand correlation: r = {corr.get('humidity', 'N/A')}\n\n"
        f"Implication: Targeted energy-efficiency programs should prioritise\n"
        f"the {top_seg} segment. Weather-linked demand suggests HVAC usage\n"
        f"is a key driver."
    )
    add_textbox(sl, f2_text, 6.8, 1.75, 6.2, 2.5, size=13)
    add_image_safe(sl, "segment_comparison.png", 6.7, 4.1, 6.2)
    add_textbox(sl, DISCLAIMER, 0.5, 7.1, 12, 0.35, size=10,
                color=RED, align=PP_ALIGN.CENTER)

    # ── Slide 4: Machine Learning ──
    sl = add_slide()
    add_rect(sl, 0, 0, 13.33, 1.1, DARK_BLUE)
    add_textbox(sl, "Machine Learning Model", 0.3, 0.15, 12, 0.8, bold=True, size=28, color=WHITE)
    ml_left = (
        f"Algorithm: Random Forest Regressor (100 trees, seed=42)\n\n"
        f"Features ({len(model_results['num_features']) + len(model_results['cat_features'])} total):\n"
        f"  Numeric: {', '.join(model_results['num_features'])}\n"
        f"  Categorical: {', '.join(model_results['cat_features'])}\n\n"
        f"Target: consumption (kWh)\n\n"
        f"Split: Chronological (80% train / 20% test)\n"
        f"  Train rows: {model_results['train_rows']:,}\n"
        f"  Test rows:  {model_results['test_rows']:,}\n"
        f"  Split at: {model_results['split_ts'][:10]}"
    )
    add_textbox(sl, ml_left, 0.4, 1.2, 6.2, 4.5, size=13)

    metrics_text = (
        f"Evaluation Metrics\n\n"
        f"{'Metric':<8}  {'Baseline':>10}  {'RandomForest':>12}\n"
        f"{'MAE':<8}  {baseline['MAE']:>10.4f}  {rf['MAE']:>12.4f}\n"
        f"{'RMSE':<8}  {baseline['RMSE']:>10.4f}  {rf['RMSE']:>12.4f}\n"
        f"{'R²':<8}  {baseline['R2']:>10.4f}  {rf['R2']:>12.4f}\n"
    )
    add_textbox(sl, metrics_text, 6.8, 1.2, 6.2, 2.0, size=14, bold=False)
    add_image_safe(sl, "actual_vs_predicted.png", 6.7, 3.3, 6.2)
    add_textbox(sl, DISCLAIMER, 0.5, 7.1, 12, 0.35, size=10,
                color=RED, align=PP_ALIGN.CENTER)

    # ── Slide 5: Recommendations ──
    sl = add_slide()
    add_rect(sl, 0, 0, 13.33, 1.1, DARK_BLUE)
    add_textbox(sl, "Business Recommendations", 0.3, 0.15, 12, 0.8, bold=True, size=28, color=WHITE)
    recs = (
        f"1. Demand-Response at Peak Hour ({peak_hour}:00)\n"
        f"   Evidence: Peak aggregate demand is {eda_results['hourly']['avg_aggregate_demand'].max():.1f} kWh "
        f"vs low of {eda_results['hourly']['avg_aggregate_demand'].min():.1f} kWh at {low_hour}:00.\n"
        f"   Action: Implement time-of-use incentives to shift flexible loads away from {peak_hour}:00.\n\n"
        f"2. Targeted Efficiency for {top_seg} Segment\n"
        f"   Evidence: {top_seg} has the highest mean consumption ({top_seg_mean:.3f} kWh/reading).\n"
        f"   Action: Offer energy audits and smart-meter upgrades to {top_seg} customers first.\n\n"
        f"3. Weather-Linked Forecasting\n"
        f"   Evidence: Temperature correlation r={corr.get('temperature','N/A')} with aggregate demand.\n"
        f"   Action: Integrate weather forecasts into grid planning to anticipate demand spikes.\n\n"
        f"4. Model Deployment\n"
        f"   The Random Forest model (R²={rf['R2']}, MAE={rf['MAE']}) outperforms the mean baseline\n"
        f"   (R²={baseline['R2']}) and can support short-term consumption forecasting.\n\n"
        f"Limitations: Synthetic data; results are illustrative only. Real deployment requires\n"
        f"validation on actual meter data and regulatory review of tariff recommendations."
    )
    add_textbox(sl, recs, 0.4, 1.2, 12.5, 5.5, size=13)
    add_textbox(sl, DISCLAIMER, 0.5, 7.1, 12, 0.35, size=10,
                color=RED, align=PP_ALIGN.CENTER)

    path = PPTX_DIR / "electricity_consumption_insights.pptx"
    prs.save(path)
    print(f"PowerPoint saved: {path} ({path.stat().st_size:,} bytes)")
    return path


if __name__ == "__main__":
    import json
    from src.analysis import run_eda
    from src.modeling import run_modeling
    eda = run_eda()
    model = run_modeling()
    create_excel(eda)
    create_pptx(eda, model)
